"""
HireAI AI Engine - Document Loader Utility.

Extracts plain text and metadata from uploaded documents (PDF, DOCX, TXT, MD, JSON).
Supports Streamlit UploadedFile objects, file paths, and raw bytes.
"""

import io
import os
from typing import Dict, Any, Union


def extract_document_info(file_obj: Union[Any, str, bytes], filename: str = None) -> Dict[str, Any]:
    """
    Extracts text and metadata from an uploaded file or file path.

    Args:
        file_obj: Streamlit UploadedFile, file-like object, bytes, or file path string.
        filename: Optional filename override if file_obj is bytes or generic stream.

    Returns:
        Dict containing:
            - text: str (extracted text)
            - filename: str
            - extension: str (.pdf, .docx, .txt, etc.)
            - char_count: int
            - word_count: int
            - estimated_tokens: int
            - page_count: int | None
            - error: str | None
    """
    # Determine filename
    if filename is None:
        if hasattr(file_obj, "name"):
            filename = file_obj.name
        elif isinstance(file_obj, str):
            filename = os.path.basename(file_obj)
        else:
            filename = "document.txt"

    extension = os.path.splitext(filename)[1].lower()

    # Read bytes
    if isinstance(file_obj, str):
        with open(file_obj, "rb") as f:
            file_bytes = f.read()
    elif hasattr(file_obj, "getvalue"):
        file_bytes = file_obj.getvalue()
    elif hasattr(file_obj, "read"):
        file_bytes = file_obj.read()
        if hasattr(file_obj, "seek"):
            file_obj.seek(0)
    elif isinstance(file_obj, bytes):
        file_bytes = file_obj
    else:
        return {
            "text": "",
            "filename": filename,
            "extension": extension,
            "char_count": 0,
            "word_count": 0,
            "estimated_tokens": 0,
            "page_count": None,
            "error": f"Unsupported file object type: {type(file_obj)}",
        }

    text = ""
    page_count = None
    error = None

    try:
        if extension == ".pdf":
            text, page_count = _extract_pdf(file_bytes)
        elif extension in [".docx", ".doc"]:
            text = _extract_docx(file_bytes)
        else:
            # Treat as text (.txt, .md, .json, .csv, etc.)
            text = _extract_plain_text(file_bytes)

        text = text.strip()
    except Exception as e:
        error = f"Failed to extract text from {filename}: {str(e)}"

    char_count = len(text)
    word_count = len(text.split()) if text else 0
    estimated_tokens = max(1, char_count // 4) if char_count > 0 else 0

    return {
        "text": text,
        "filename": filename,
        "extension": extension,
        "char_count": char_count,
        "word_count": word_count,
        "estimated_tokens": estimated_tokens,
        "page_count": page_count,
        "error": error,
    }


def _extract_pdf(file_bytes: bytes) -> tuple[str, int]:
    """Extracts text and page count from a PDF file using pypdf."""
    from pypdf import PdfReader

    stream = io.BytesIO(file_bytes)
    reader = PdfReader(stream)
    pages_text = []

    for page_idx, page in enumerate(reader.pages):
        page_str = page.extract_text() or ""
        if page_str.strip():
            pages_text.append(page_str.strip())

    return "\n\n".join(pages_text), len(reader.pages)


def _extract_docx(file_bytes: bytes) -> str:
    """Extracts text from a DOCX file using python-docx."""
    from docx import Document

    stream = io.BytesIO(file_bytes)
    doc = Document(stream)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    # Also extract text from tables if any
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                paragraphs.append(" | ".join(row_text))

    return "\n".join(paragraphs)


def _extract_plain_text(file_bytes: bytes) -> str:
    """Decodes plain text with automatic fallback for different encodings."""
    for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
        try:
            return file_bytes.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return file_bytes.decode("utf-8", errors="replace")
